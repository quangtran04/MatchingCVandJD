from __future__ import annotations

from io import BytesIO
from pathlib import Path


class TextExtractionError(Exception):
    def __init__(self, code: str, message: str, status_code: int = 422) -> None:
        self.code = code
        self.message = message
        self.status_code = status_code
        super().__init__(message)


SUPPORTED_SUFFIXES = {".txt", ".pdf", ".docx"}


def extract_text(payload: bytes, filename: str) -> tuple[str, str]:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise TextExtractionError(
            "UNSUPPORTED_FILE_TYPE",
            "Chỉ hỗ trợ TXT, PDF có text hoặc Word (.docx).",
            415,
        )
    if not payload:
        raise TextExtractionError("EMPTY_FILE", "Tệp tải lên đang trống.")

    if suffix == ".txt":
        return _extract_txt(payload), "plain_text"
    if suffix == ".docx":
        return _extract_docx(payload), "docx"
    if suffix == ".pdf":
        return _extract_pdf(payload)

    raise TextExtractionError("UNSUPPORTED_FILE_TYPE", "Định dạng không được hỗ trợ.")


def _extract_txt(payload: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "utf-16", "cp1258", "latin-1"):
        try:
            text = payload.decode(encoding).strip()
            if text:
                return text
        except UnicodeDecodeError:
            continue
    raise TextExtractionError("INVALID_TEXT_ENCODING", "Không thể đọc mã hoá của tệp TXT.")


def _extract_docx(payload: bytes) -> str:
    try:
        from docx import Document
    except ImportError as error:
        raise TextExtractionError(
            "PARSER_NOT_INSTALLED", "Thiếu thư viện python-docx trên máy chủ.", 503
        ) from error

    try:
        document = Document(BytesIO(payload))
        parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text.strip() for cell in row.cells if cell.text.strip())
    except Exception as error:
        raise TextExtractionError("INVALID_DOCX", "Không thể đọc tệp Word (.docx).") from error

    text = "\n".join(parts).strip()
    if not text:
        raise TextExtractionError("EMPTY_DOCX", "Tệp Word không chứa nội dung văn bản.")
    return text


def _extract_pdf(payload: bytes) -> tuple[str, str]:
    text = ""
    method = ""

    try:
        import fitz
        document = fitz.open(stream=payload, filetype="pdf")
        text = "\n".join(page.get_text("text") for page in document).strip()
        method = "pdf_text"
    except ImportError:
        try:
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(payload))
            text = "\n".join(page.extract_text() or "" for page in reader.pages).strip()
            method = "pdf_text"
        except ImportError as err:
            raise TextExtractionError(
                "PARSER_NOT_INSTALLED",
                "Thiếu thư viện đọc PDF trên máy chủ.",
                503,
            ) from err
        except Exception as err:
            raise TextExtractionError("INVALID_PDF", "Không thể phân tích nội dung PDF.") from err
    except Exception as error:
        raise TextExtractionError("INVALID_PDF", "Không thể đọc tệp PDF.") from error

    if not text:
        raise TextExtractionError(
            "PDF_SCAN_NOT_SUPPORTED",
            "PDF không có lớp text. Phiên bản này chưa hỗ trợ OCR.",
        )

    return text, method